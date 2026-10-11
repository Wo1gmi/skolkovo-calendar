"""LMS pilot: exact assignment deadlines plus a review of every course file.

File snippets are candidates, never automatically invented deadlines.
Private materials are not printed or uploaded as public workflow artifacts.
"""
import hashlib
import io
import json
import os
import re
from datetime import datetime, timedelta
from urllib.parse import urljoin, urlparse
from zoneinfo import ZoneInfo

import requests
from bs4 import BeautifulSoup

BASE = 'https://l.skolkovo.ru'
TZ = ZoneInfo('Europe/Moscow')
MARKER = 'skolkovo_lms_assignments_v1'
MONTHS = dict(zip(('января февраля марта апреля мая июня июля августа сентября октября ноября декабря').split(), range(1, 13)))


def parse_due(text):
    match = re.search(r'(?:Срок сдачи|Due date)\s*:\s*[^\n]*?(\d{1,2})\s+([а-я]+)\s+(\d{4}),?\s+(\d{1,2}):(\d{2})', text, re.I)
    if not match or match[2].lower() not in MONTHS:
        return None
    return datetime(int(match[3]), MONTHS[match[2].lower()], int(match[1]), int(match[4]), int(match[5]), tzinfo=TZ)


def safe_url(url):
    parsed = urlparse(url)
    if parsed.scheme != 'https' or parsed.hostname != 'l.skolkovo.ru':
        raise RuntimeError('Unexpected LMS destination')
    return url


class LMS:
    def __init__(self):
        self.session = requests.Session()

    def page(self, url):
        response = self.session.get(safe_url(url), timeout=45)
        response.raise_for_status()
        if '/login/' in response.url:
            raise RuntimeError('LMS session expired')
        return BeautifulSoup(response.text, 'html.parser')

    def login(self):
        url = BASE + '/login/index.php'
        response = self.session.get(url, timeout=45)
        response.raise_for_status()
        soup = BeautifulSoup(response.text, 'html.parser')
        form = next((f for f in soup.find_all('form') if f.find('input', {'type': 'password'})), None)
        if form is None:
            raise RuntimeError('LMS login form not found')
        data = {e['name']: e.get('value', '') for e in form.select('input[name]') if e.get('type') == 'hidden'}
        password = form.find('input', {'type': 'password'})
        username = next((e for e in form.select('input[name]') if e.get('type', 'text') in ('text', 'email')), None)
        if username is None:
            raise RuntimeError('LMS username field not found')
        data[username['name']] = os.environ['SKOLKOVO_EMAIL']
        data[password['name']] = os.environ['SKOLKOVO_PASSWORD']
        response = self.session.post(safe_url(urljoin(url, form.get('action') or url)), data=data, timeout=45)
        response.raise_for_status()
        self.page(BASE + '/my/')

    def file(self, url):
        with self.session.get(safe_url(url), timeout=60, stream=True) as response:
            response.raise_for_status()
            if '/login/' in response.url or 'text/html' in response.headers.get('Content-Type', ''):
                raise RuntimeError('File download returned HTML')
            chunks, size = [], 0
            for chunk in response.iter_content(65536):
                size += len(chunk)
                if size > 80 * 1024 * 1024:
                    raise RuntimeError('File exceeds 80 MB limit')
                chunks.append(chunk)
            return b''.join(chunks)


def file_text(name, data):
    ext = name.lower().split('.')[-1]
    stream = io.BytesIO(data)
    if ext == 'pdf':
        from pypdf import PdfReader
        pages = [p.extract_text() or '' for p in PdfReader(stream).pages]
        if not any(p.strip() for p in pages):
            raise RuntimeError('PDF requires OCR')
        return '\n'.join(pages)
    if ext == 'docx':
        from docx import Document
        doc = Document(stream)
        return '\n'.join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
    if ext == 'pptx':
        from pptx import Presentation
        return '\n'.join(shape.text for slide in Presentation(stream).slides for shape in slide.shapes if shape.has_text_frame)
    if ext == 'xlsx':
        from openpyxl import load_workbook
        book = load_workbook(stream, read_only=True, data_only=False)
        return '\n'.join(' | '.join(str(v) for v in row if v is not None) for sheet in book for row in sheet.values)
    if ext in ('txt', 'csv'):
        return data.decode('utf-8-sig')
    raise RuntimeError('Unsupported file format: ' + ext)


def candidates(text):
    # Report surrounding instructions, not every date or occurrence of "task".
    pattern = re.compile(r'домашн\w*\s+задан|дедлайн|срок\w*\s+сдач|к\s+заключительн\w*\s+занят|выполн\w*\s+задан|во время чтения|подготов\w*\s+(?:презентац|доклад)|deadline|homework|due date', re.I)
    chunks = []
    lines = text.splitlines()
    covered = -1
    for index, line in enumerate(lines):
        if index > covered and pattern.search(line):
            covered = min(len(lines), index + 7)
            chunks.append('\n'.join(lines[max(0, index-1):covered]))
    return chunks


def telegram(text):
    for offset in range(0, len(text), 3500):
        response = requests.post('https://api.telegram.org/bot' + os.environ['TELEGRAM_BOT_TOKEN'] + '/sendMessage', json={
            'chat_id': os.environ['TELEGRAM_CHAT_ID'], 'text': text[offset:offset+3500], 'disable_web_page_preview': True}, timeout=30)
        if response.status_code != 200 or not response.json().get('ok'):
            raise RuntimeError('Telegram delivery failed')


def collect(client, course_ids):
    assignments, reviews, failures = [], [], []
    for course_id in course_ids:
        page = client.page(BASE + '/course/view.php?id=' + str(course_id))
        title = page.title.get_text().split('|')[0].removeprefix('Курс:').strip()
        resources = {}
        for a in page.select('a[href]'):
            url = urljoin(BASE, a['href'])
            if re.search(r'/mod/(?:resource|assign|folder|page)/view.php\?id=\d+', url):
                resources[url.split('&')[0]] = a.get_text(' ', strip=True)
        files_seen = set()
        for url, label in resources.items():
            material = client.page(url)
            region = material.select_one('#region-main') or material
            text = region.get_text('\n', strip=True)
            if '/mod/assign/' in url:
                due = parse_due(text)
                if due:
                    assignments.append({'id': re.search(r'id=(\d+)', url)[1], 'title': (material.find('h1').get_text(strip=True) if material.find('h1') else label), 'course': title, 'due': due.isoformat(), 'url': url})
                else:
                    reviews.append({'course': title, 'name': label, 'url': url, 'snippets': ['Задание LMS: точный срок не найден; требуется проверка.']})
            elif candidates(text):
                reviews.append({'course': title, 'name': label, 'url': url, 'snippets': candidates(text)})
            for link in region.select('a[href*="/pluginfile.php/"]'):
                file_url = urljoin(BASE, link['href'])
                if file_url in files_seen:
                    continue
                files_seen.add(file_url)
                name = link.get_text(strip=True)
                try:
                    snippets = candidates(file_text(name, client.file(file_url)))
                    if snippets:
                        reviews.append({'course': title, 'name': name, 'url': url, 'snippets': snippets})
                except Exception as error:
                    failures.append({'course': title, 'name': name, 'url': url, 'reason': str(error) if isinstance(error, RuntimeError) else type(error).__name__})
        print('Course scanned:', course_id, 'files:', len(files_seen))
    return assignments, reviews, failures


def future_assignments(assignments, now=None):
    now = now or datetime.now(TZ)
    return [a for a in assignments if datetime.fromisoformat(a['due']) > now]


def assignment_body(assignment):
    due = datetime.fromisoformat(assignment['due']).astimezone(TZ)
    day = due.date() - timedelta(days=1)
    return {
        'summary': 'Сдать: ' + assignment['course'] + ' — ' + assignment['title'],
        'colorId': '11',
        'description': 'Срок сдачи: ' + due.strftime('%d.%m.%Y %H:%M') +
                       ' (Москва).\\nСобытие размещено за день до дедлайна.\\n' + assignment['url'],
        'start': {'date': day.isoformat()},
        'end': {'date': due.date().isoformat()},
        'reminders': {'useDefault': False, 'overrides': [
            {'method': 'popup', 'minutes': 1440}]},
    }


def sync_exact(assignments):
    from google.oauth2 import service_account
    from googleapiclient.discovery import build
    credentials = service_account.Credentials.from_service_account_info(json.loads(os.environ['GOOGLE_SERVICE_ACCOUNT_JSON']), scopes=['https://www.googleapis.com/auth/calendar.events'])
    api = build('calendar', 'v3', credentials=credentials, cache_discovery=False).events()
    calendar = os.environ.get('LMS_CALENDAR_ID') or os.environ['GOOGLE_CALENDAR_ID']
    current, page_token = {}, None
    while True:
        result = api.list(calendarId=calendar, privateExtendedProperty='sync_source=' + MARKER, maxResults=2500, pageToken=page_token).execute()
        for event in result.get('items', []):
            current[event.get('extendedProperties', {}).get('private', {}).get('source_id')] = event
        page_token = result.get('nextPageToken')
        if not page_token:
            break
    now = datetime.now(TZ)
    for assignment in assignments:
        due = datetime.fromisoformat(assignment['due'])
        if due <= now:
            continue
        old = current.get(assignment['id'])
        private = dict((old or {}).get('extendedProperties', {}).get('private', {}))
        private.update(sync_source=MARKER, source_id=assignment['id'])
        body = assignment_body(assignment)
        if old:
            # Clear timed fields when migrating existing events to all-day dates.
            for field in ('start', 'end'):
                if old.get(field, {}).get('dateTime'):
                    body[field].update(dateTime=None, timeZone=None)
        digest = hashlib.sha256(json.dumps(body, sort_keys=True).encode()).hexdigest()
        changed = private.get('content_hash') != digest
        private['content_hash'] = digest
        body['extendedProperties'] = {'private': private}
        if old is None:
            body['colorId'] = '11'
            # Deterministic event IDs avoid duplicates after interrupted runs.
            event_id = hashlib.sha256((MARKER + assignment['id']).encode()).hexdigest()
            body['id'] = event_id
            api.insert(calendarId=calendar, body=body).execute()
        elif changed:
            event_id = old['id']
            api.patch(calendarId=calendar, eventId=event_id, body=body).execute()
        else:
            event_id = old['id']
        if changed:
            telegram(('📝 Новое задание' if old is None else '🔄 Изменён дедлайн/задание') + '\n' + body['summary'] + '\n' + due.strftime('%d.%m.%Y %H:%M') + '\n' + assignment['url'])
        # One reminder per threshold and deadline, with state in the private calendar.
        hours = (due - now).total_seconds() / 3600
        bucket = '24' if hours <= 24 else '72' if hours <= 72 else None
        key = 'reminded_' + bucket if bucket else None
        if key and private.get(key) != due.isoformat():
            telegram('⏳ До сдачи осталось примерно ' + str(max(1, round(hours))) + ' ч.\n' + body['summary'] + '\n' + assignment['url'])
            private[key] = due.isoformat()
            api.patch(calendarId=calendar, eventId=event_id, body={'extendedProperties': {'private': private}}).execute()


def main():
    config = json.loads(open('lms_courses.json', encoding='utf-8').read())
    client = LMS()
    client.login()
    assignments, reviews, failures = collect(client, config['course_ids'])
    print('Exact deadlines:', len(assignments), 'file candidates:', len(reviews), 'unread files:', len(failures))
    assignments = future_assignments(assignments)
    print('Future exact deadlines:', len(assignments))
    if os.environ.get('LMS_APPLY') == 'true':
        sync_exact(assignments)
    elif os.environ.get('LMS_REPORT') == 'true':
        for assignment in assignments:
            due = datetime.fromisoformat(assignment['due'])
            telegram('📝 Задание: ' + assignment['course'] + ' — ' + assignment['title'] +
                     '\nСдать до: ' + due.strftime('%d.%m.%Y %H:%M') +
                     '\n' + assignment['url'])
    # Undated candidates and download failures stay out of Telegram.


if __name__ == '__main__':
    try:
        main()
    except Exception as error:
        print('LMS failed:', str(error) if isinstance(error, RuntimeError) else type(error).__name__)
        raise SystemExit(1)
